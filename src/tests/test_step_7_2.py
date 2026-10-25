"""
Booklet_AI Desktop — STEP-7.2 Unit & Integration Tests (DOCX Generation CLI & Report Builder)

Verifies:
1. Extraction of a real DOCX fixture into a BookletDocument with hash verification.
2. Rule-based classification preserving Zero-Mutation guarantee.
3. build_report() accurately aggregates:
   - Source paragraph, table, equation counts.
   - Breakdown of blocks per block_type.
   - Legacy equation count.
   - Math formula count and first 3 formulas sample.
   - Actual page count and per-page diagnostics.
   - Stage timings and warnings list.
4. CLI run_pipeline with --build-only and --no-ai produces valid payload and report.json.
5. Error handling exits gracefully on invalid file paths.
"""

import json
import sys
from pathlib import Path
from unittest.mock import MagicMock, patch

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import pytest
from docx import Document as DocxDocument

from app.core.models import BookletDocument, ContentBlock, PageData, TableData
from app.core.text_integrity import compute_canonical_hash
from tools.generate_from_docx import (
    build_report,
    classify_document,
    extract_document,
    format_report_text,
    run_pipeline,
)


@pytest.fixture
def sample_docx_fixture(tmp_path: Path) -> Path:
    """Create a minimal real Word document with paragraphs, table, and question."""
    doc_path = tmp_path / "sample_curriculum.docx"
    doc = DocxDocument()

    # 1. Heading
    doc.add_heading("الفصل الأول: التفاضل والتكامل", level=1)

    # 2. Arabic body paragraph
    doc.add_paragraph(
        "يُعد التفاضل أحد أهم فروع الرياضيات الحديثة، ويُستخدم لحساب معدلات التغير "
        "والمماسات لمنحنيات الدوال المختلفة في التطبيقات الهندسية والفيزيائية."
    )

    # 3. Ministry Question
    doc.add_paragraph(
        "س وزاري: جد معادلة المماس لمنحني الدالة f(x) = x² + 3 عند النقطة (1, 4)."
    )

    # 4. List Item
    doc.add_paragraph("- الخاصية الأولى: مشتقة الثابت تساوي صفراً.")

    # 5. Table (2 rows, 2 columns)
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "الدالة f(x)"
    tbl.cell(0, 1).text = "المشتقة f'(x)"
    tbl.cell(1, 0).text = "c"
    tbl.cell(1, 1).text = "0"

    doc.save(str(doc_path))
    return doc_path


class TestDocxPipelineExtractionAndClassification:
    def test_extract_document_fixture(self, sample_docx_fixture: Path):
        """Extract content from real DOCX fixture and verify source counts and integrity."""
        booklet_doc, source_counts, elapsed = extract_document(sample_docx_fixture)

        assert isinstance(booklet_doc, BookletDocument)
        assert len(booklet_doc.pages) == 1
        page = booklet_doc.pages[0]

        # Verify source counts
        assert source_counts["paragraphs"] >= 4
        assert source_counts["tables"] == 1
        assert source_counts["equations"] == 0
        assert elapsed >= 0.0

        # Verify extracted blocks have matching canonical hashes
        assert len(page.content_blocks) >= 4
        for b in page.content_blocks:
            assert b.content_hash == compute_canonical_hash(b.raw_content)

    def test_classify_document_offline(self, sample_docx_fixture: Path):
        """Classify document blocks with RuleBasedClassifier under --no-ai."""
        booklet_doc, _, _ = extract_document(sample_docx_fixture)
        classified_doc, warnings, elapsed = classify_document(booklet_doc, no_ai=True)

        assert elapsed >= 0.0
        assert isinstance(warnings, list)
        page = classified_doc.pages[0]

        # Check that classification mapped expected types
        types = [b.block_type for b in page.content_blocks]
        assert "chapter_header" in types
        assert "arabic_text" in types
        assert "ministry_question" in types
        assert "list_item" in types

        # Zero-Mutation guarantee
        for b in page.content_blocks:
            assert b.content_hash == compute_canonical_hash(b.raw_content)


class TestReportBuilder:
    def test_build_report_structure(self):
        """Verify build_report generates all required fields."""
        blk1 = ContentBlock(
            block_id="b1",
            block_type="chapter_header",
            raw_content="الفصل الأول",
            content_hash=compute_canonical_hash("الفصل الأول"),
        )
        blk2 = ContentBlock(
            block_id="b2",
            block_type="math_formula",
            raw_content="x^2 + y^2 = r^2",
            content_hash=compute_canonical_hash("x^2 + y^2 = r^2"),
        )
        blk3 = ContentBlock(
            block_id="b3",
            block_type="math_formula",
            raw_content="E = mc^2",
            content_hash=compute_canonical_hash("E = mc^2"),
        )
        blk4 = ContentBlock(
            block_id="b4",
            block_type="legacy_equation_placeholder",
            raw_content="[OLE: MathType 5.0]",
            content_hash=compute_canonical_hash("[OLE: MathType 5.0]"),
        )

        page = PageData(
            page_number=1,
            width=595.28,
            height=841.89,
            content_blocks=[blk1, blk2, blk3, blk4],
            tables=[],
        )
        booklet_doc = BookletDocument(
            doc_id="report_test_doc",
            file_path="C:/test/sample.docx",
            file_type="docx",
            pages=[page],
        )

        source_counts = {
            "paragraphs": 5,
            "tables": 1,
            "equations": 3,
            "omml_math": 2,
            "legacy_equations": 1,
        }
        timings = {
            "extraction_seconds": 0.12,
            "classification_seconds": 0.05,
            "payload_build_seconds": 0.08,
            "indesign_generation_seconds": 1.25,
            "total_seconds": 1.50,
        }
        indesign_result = {
            "actual_page_count": 4,
            "imposed_page_count": 4,
            "page_count": 4,
            "pages": [
                {
                    "page_number": 1,
                    "side": "left_hand",
                    "geometric_bounds": [62.69, 42.52, 779.20, 532.92],
                    "table_width": None,
                }
            ],
            "warnings": ["Font substituted"],
        }
        pipeline_warnings = ["AI key not set"]

        report = build_report(
            booklet_doc=booklet_doc,
            source_counts=source_counts,
            timings=timings,
            indesign_result=indesign_result,
            pipeline_warnings=pipeline_warnings,
            output_files={"report_json": "report.json"},
        )

        # 1. Source counts
        assert report["source_counts"]["paragraphs"] == 5
        assert report["source_counts"]["tables"] == 1
        assert report["source_counts"]["equations"] == 3

        # 2. Blocks per block_type
        assert report["blocks_per_type"]["chapter_header"] == 1
        assert report["blocks_per_type"]["math_formula"] == 2
        assert report["blocks_per_type"]["legacy_equation_placeholder"] == 1

        # 3. Math & Legacy equation counts and first 3 sample
        assert report["legacy_equation_count"] == 1
        assert report["math_formula_count"] == 2
        assert len(report["first_3_math_formulas"]) == 2
        assert report["first_3_math_formulas"][0] == "x^2 + y^2 = r^2"
        assert report["first_3_math_formulas"][1] == "E = mc^2"

        # 4. Actual page count & per-page diagnostics
        assert report["actual_page_count"] == 4
        assert report["imposed_page_count"] == 4
        assert len(report["per_page_diagnostics"]) == 1
        assert report["per_page_diagnostics"][0]["side"] == "left_hand"

        # 5. Warnings and timings
        assert "AI key not set" in report["warnings"]
        assert "Font substituted" in report["warnings"]
        assert report["timings"]["total_seconds"] == 1.50

    def test_format_report_text(self):
        """Verify format_report_text produces readable multi-line summary."""
        report = {
            "doc_id": "d1",
            "source_file": "sample.docx",
            "source_counts": {"paragraphs": 10, "tables": 2, "equations": 3, "omml_math": 2, "legacy_equations": 1},
            "blocks_per_type": {"arabic_text": 8, "math_formula": 2},
            "total_blocks": 10,
            "legacy_equation_count": 1,
            "math_formula_count": 2,
            "first_3_math_formulas": ["a + b = c"],
            "actual_page_count": 4,
            "imposed_page_count": 4,
            "final_page_count": 4,
            "per_page_diagnostics": [{"page_number": 1, "side": "left_hand", "geometric_bounds": [1, 2, 3, 4], "table_width": None}],
            "warnings": ["Sample warning"],
            "timings": {"extraction_seconds": 0.1, "total_seconds": 1.0},
            "output_files": {"report_json": "report.json"},
        }

        text = format_report_text(report)
        assert "Booklet_AI — STEP-7.2 Real DOCX Measurement Report" in text
        assert "Source Document Counts" in text
        assert "Paragraphs:       10" in text
        assert "Math Formula Blocks:          2" in text
        assert "a + b = c" in text
        assert "Actual Content Pages:  4" in text
        assert "Sample warning" in text


class TestPipelineExecution:
    def test_run_pipeline_build_only_produces_valid_payload_and_report(
        self, sample_docx_fixture: Path, tmp_path: Path
    ):
        """End-to-end test of run_pipeline in --build-only mode."""
        out_dir = tmp_path / "out_build"

        exit_code = run_pipeline(
            docx_file=sample_docx_fixture,
            output_dir=out_dir,
            binding="saddle_stitch",
            footer_text="اختبار الفصل",
            no_ai=True,
            build_only=True,
        )

        assert exit_code == 0
        report_file = out_dir / "report.json"
        assert report_file.exists()

        report_data = json.loads(report_file.read_text(encoding="utf-8"))
        assert report_data["source_counts"]["paragraphs"] >= 4
        assert report_data["source_counts"]["tables"] == 1
        assert "chapter_header" in report_data["blocks_per_type"]

        # Check payload was also created
        payload_files = list(out_dir.glob("*_payload.json"))
        assert len(payload_files) == 1
        payload_data = json.loads(payload_files[0].read_text(encoding="utf-8"))
        assert payload_data["footer_text"] == "اختبار الفصل"

    def test_run_pipeline_file_not_found(self, tmp_path: Path):
        """Verify pipeline handles missing file gracefully."""
        missing_file = tmp_path / "does_not_exist.docx"
        exit_code = run_pipeline(
            docx_file=missing_file,
            output_dir=tmp_path / "out",
            no_ai=True,
            build_only=True,
        )
        assert exit_code == 1
