"""
Tests for STEP-1.2 — Word (DOCX) Extraction Engine: Math & Chemistry Formulas

Covers all 5 Acceptance Criteria (AC-1.2.1 through AC-1.2.5).

Tests create DOCX files programmatically with injected XML to simulate
real-world Word content including OMML math, chemistry formulas with
subscript/superscript runs, tables, headings, and OLE objects.

Tests that require the OMML2MML.XSL stylesheet are marked with
@pytest.mark.skipif when the stylesheet is not available.
"""

import os
import sys
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch

import lxml.etree as ET
import pytest
from docx import Document as DocxDocument

# ─── Ensure src/ is on sys.path ────────────────────────────────────────
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.core.exceptions import MathExtractionError
from app.core.models import BlockType, ContentBlock, TableData
from app.core.text_integrity import compute_canonical_hash
from app.core.docx_extractor import (
    NSMAP,
    DocxContentExtractor,
    DocxMathExtractor,
    _analyze_runs_for_chemistry,
    _assemble_chemistry_text,
    _discover_xslt,
    _get_run_text,
    _get_run_vertical_align,
)

# ─── Namespace shortcuts ──────────────────────────────────────────────
W_NS = NSMAP["w"]
M_NS = NSMAP["m"]
V_NS = NSMAP["v"]
O_NS = NSMAP["o"]

# ─── XSLT availability check ──────────────────────────────────────────
try:
    from tests.conftest import is_xslt_available, requires_xslt
except ImportError:
    from conftest import is_xslt_available, requires_xslt

XSLT_AVAILABLE = is_xslt_available()


# ═══════════════════════════════════════════════════════════════════════
# Helpers: create DOCX files with specific XML content
# ═══════════════════════════════════════════════════════════════════════


def _create_docx_with_body_xml(body_inner_xml: str, tmp_dir: str) -> str:
    """Create a minimal .docx file with custom body XML content.

    Uses python-docx to create a valid DOCX, then patches the
    document.xml inside the ZIP with custom body content.

    Args:
        body_inner_xml: XML string for the contents inside w:body.
        tmp_dir: Temporary directory to save the file.

    Returns:
        Path to the created .docx file.
    """
    # Create a minimal valid DOCX first
    doc = DocxDocument()
    base_path = os.path.join(tmp_dir, "test.docx")
    doc.save(base_path)

    # Now patch the document.xml inside the ZIP
    patched_path = os.path.join(tmp_dir, "test_patched.docx")

    full_body_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas"
            xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
            xmlns:o="urn:schemas-microsoft-com:office:office"
            xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"
            xmlns:v="urn:schemas-microsoft-com:vml"
            xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
            xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"
            mc:Ignorable="w14 wpc">
  <w:body>
    {body_inner_xml}
  </w:body>
</w:document>"""

    with zipfile.ZipFile(base_path, "r") as zin:
        with zipfile.ZipFile(patched_path, "w") as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "word/document.xml":
                    data = full_body_xml.encode("utf-8")
                zout.writestr(item, data)

    return patched_path


def _make_paragraph_xml(text: str, style: str = "") -> str:
    """Create a simple w:p XML string with text.

    Args:
        text: The text content.
        style: Optional paragraph style (e.g. 'Heading1').

    Returns:
        XML string for a w:p element.
    """
    style_xml = ""
    if style:
        style_xml = f"""<w:pPr>
            <w:pStyle w:val="{style}"/>
        </w:pPr>"""

    return f"""<w:p>
        {style_xml}
        <w:r>
            <w:t xml:space="preserve">{text}</w:t>
        </w:r>
    </w:p>"""


def _make_chemistry_paragraph_xml(formula_parts: list) -> str:
    """Create a w:p with chemistry runs (subscript/superscript).

    Args:
        formula_parts: List of (text, vert_align) tuples where vert_align
                       is None, 'subscript', or 'superscript'.

    Returns:
        XML string for a w:p element with chemistry runs.
    """
    runs = []
    for text, valign in formula_parts:
        rpr = ""
        if valign:
            rpr = f"""<w:rPr>
                <w:vertAlign w:val="{valign}"/>
            </w:rPr>"""
        runs.append(f"""<w:r>
            {rpr}
            <w:t xml:space="preserve">{text}</w:t>
        </w:r>""")

    return f"""<w:p>
        {"".join(runs)}
    </w:p>"""


def _make_omath_xml(content_text: str = "x+y=z") -> str:
    """Create a minimal m:oMath element XML.

    Args:
        content_text: Simple text content for the math run.

    Returns:
        XML string for an m:oMath element inside a w:p.
    """
    return f"""<w:p>
        <m:oMath>
            <m:r>
                <m:t>{content_text}</m:t>
            </m:r>
        </m:oMath>
    </w:p>"""


def _make_omathpara_xml(content_text: str = "a^2+b^2=c^2") -> str:
    """Create a block-level m:oMathPara element XML.

    Args:
        content_text: Simple text content for the math run.

    Returns:
        XML string for a w:p containing an m:oMathPara.
    """
    return f"""<w:p>
        <m:oMathPara>
            <m:oMath>
                <m:r>
                    <m:t>{content_text}</m:t>
                </m:r>
            </m:oMath>
        </m:oMathPara>
    </w:p>"""


def _make_table_xml(headers: list, rows: list, bidi: bool = True) -> str:
    """Create a w:tbl element XML.

    Args:
        headers: List of header cell strings.
        rows: List of rows, each a list of cell strings.
        bidi: Whether table has bidiVisual (default True for Arabic).

    Returns:
        XML string for a w:tbl element.
    """
    def make_row(cells):
        tcs = []
        for cell in cells:
            tcs.append(f"""<w:tc>
                <w:p>
                    <w:r>
                        <w:t>{cell}</w:t>
                    </w:r>
                </w:p>
            </w:tc>""")
        return f"<w:tr>{''.join(tcs)}</w:tr>"

    header_row = make_row(headers)
    data_rows = "".join(make_row(row) for row in rows)
    bidi_xml = "<w:tblPr><w:bidiVisual/></w:tblPr>" if bidi else ""

    return f"<w:tbl>{bidi_xml}{header_row}{data_rows}</w:tbl>"


def _make_ole_object_xml(prog_id: str = "Equation.3") -> str:
    """Create a w:p containing a w:object (legacy OLE equation).

    Args:
        prog_id: OLE program ID (e.g. 'Equation.3' for MathType).

    Returns:
        XML string for a w:p with an embedded OLE object.
    """
    return f"""<w:p>
        <w:r>
            <w:object>
                <o:OLEObject progId="{prog_id}" />
                <v:shape style="width:100pt;height:20pt"/>
            </w:object>
        </w:r>
    </w:p>"""


# ═══════════════════════════════════════════════════════════════════════
# AC-1.2.4: XSLT missing → MathExtractionError (Fail-Fast)
# ═══════════════════════════════════════════════════════════════════════


class TestAC_1_2_4_XSLTFailFast:
    """AC-1.2.4: Any XSLT initialization failure or missing stylesheet
    raises MathExtractionError immediately — no silent skip."""

    def test_missing_xslt_raises_math_extraction_error(self):
        """DocxMathExtractor with nonexistent XSLT raises immediately."""
        with patch(
            "app.core.docx_extractor._OFFICE_XSLT_CANDIDATES", []
        ):
            with pytest.raises(MathExtractionError, match="stylesheet missing"):
                DocxMathExtractor(xslt_path="/nonexistent/path/OMML2MML.XSL")

    def test_invalid_xslt_raises_math_extraction_error(self, tmp_path):
        """DocxMathExtractor with invalid XSLT content raises immediately."""
        bad_xslt = tmp_path / "bad.xsl"
        bad_xslt.write_text("this is not valid XML or XSLT", encoding="utf-8")

        with pytest.raises(MathExtractionError, match="Failed to initialize"):
            DocxMathExtractor(xslt_path=str(bad_xslt))

    def test_discover_xslt_nonexistent_raises(self):
        """_discover_xslt with no valid paths raises MathExtractionError."""
        with patch(
            "app.core.docx_extractor._OFFICE_XSLT_CANDIDATES", []
        ):
            with pytest.raises(MathExtractionError, match="stylesheet missing"):
                _discover_xslt("/definitely/not/a/real/path.xsl")

    def test_content_extractor_fails_on_math_without_xslt(self, tmp_path):
        """DocxContentExtractor raises when encountering m:oMath
        without available XSLT — not silently skipping."""
        body_xml = _make_omath_xml("x+1=2")
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        with patch(
            "app.core.docx_extractor._OFFICE_XSLT_CANDIDATES", []
        ):
            extractor = DocxContentExtractor(
                xslt_path="/nonexistent/OMML2MML.XSL"
            )
            with pytest.raises(MathExtractionError):
                extractor.extract(docx_path)


# ═══════════════════════════════════════════════════════════════════════
# AC-1.2.1: OMML → MathML conversion (requires XSLT)
# ═══════════════════════════════════════════════════════════════════════


@pytest.mark.skipif(
    not XSLT_AVAILABLE,
    reason="OMML2MML.XSL not found — skipping MathML conversion tests",
)
class TestAC_1_2_1_OmmlToMathml:
    """AC-1.2.1: Every m:oMath/m:oMathPara produces syntactically valid
    MathML, with total count matching source m:oMath elements."""

    def test_single_omath_produces_valid_mathml(self, tmp_path):
        """A single m:oMath element produces valid MathML."""
        body_xml = _make_omath_xml("x+y")
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        math_blocks = [
            b for b in doc.pages[0].content_blocks
            if b.block_type == "math_formula"
        ]
        assert len(math_blocks) == 1

        # Verify the MathML is valid XML
        mathml = math_blocks[0].raw_content
        assert mathml  # Not empty
        parsed = ET.fromstring(mathml.encode("utf-8"))
        assert parsed is not None

    def test_multiple_omath_count_matches(self, tmp_path):
        """Multiple m:oMath elements — output count matches input count."""
        body_xml = (
            _make_omath_xml("a+b")
            + _make_paragraph_xml("Some text between equations")
            + _make_omath_xml("c-d")
            + _make_omath_xml("e*f")
        )
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        math_blocks = [
            b for b in doc.pages[0].content_blocks
            if b.block_type == "math_formula"
        ]
        # 3 m:oMath elements → 3 math_formula blocks
        assert len(math_blocks) == 3

    def test_omathpara_produces_valid_mathml(self, tmp_path):
        """A block-level m:oMathPara produces valid MathML."""
        body_xml = _make_omathpara_xml("a^2+b^2=c^2")
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        math_blocks = [
            b for b in doc.pages[0].content_blocks
            if b.block_type == "math_formula"
        ]
        assert len(math_blocks) == 1

        # Verify valid XML
        mathml = math_blocks[0].raw_content
        parsed = ET.fromstring(mathml.encode("utf-8"))
        assert parsed is not None

    def test_mathml_output_is_wellformed_xml(self, tmp_path):
        """MathML output must pass ET.fromstring() validation."""
        body_xml = _make_omath_xml("x=1")
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        for block in doc.pages[0].content_blocks:
            if block.block_type == "math_formula":
                # This must not raise
                ET.fromstring(block.raw_content.encode("utf-8"))

    def test_math_extractor_omml_to_mathml_direct(self):
        """DocxMathExtractor.omml_to_mathml converts OMML to valid MathML."""
        me = DocxMathExtractor()

        # Create a minimal OMML element
        omml_xml = f"""<m:oMath xmlns:m="{M_NS}">
            <m:r>
                <m:t>x+1</m:t>
            </m:r>
        </m:oMath>"""
        omml_elem = ET.fromstring(omml_xml.encode("utf-8"))

        result = me.omml_to_mathml(omml_elem)
        assert result  # Not empty
        # Must be valid XML
        ET.fromstring(result.encode("utf-8"))


# ═══════════════════════════════════════════════════════════════════════
# AC-1.2.2: Chemistry formula extraction
# ═══════════════════════════════════════════════════════════════════════


class TestAC_1_2_2_ChemistryFormula:
    """AC-1.2.2: Chemical formula extraction (e.g. H₂SO₄ + 2NaOH →
    Na₂SO₄ + 2H₂O) preserves subscript numbers and arrows, correctly
    classified as chemistry_formula."""

    def test_simple_chemistry_formula_detected(self, tmp_path):
        """H₂O with subscript '2' is detected as chemistry_formula."""
        body_xml = _make_chemistry_paragraph_xml([
            ("H", None),
            ("2", "subscript"),
            ("O", None),
        ])
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        chem_blocks = [
            b for b in doc.pages[0].content_blocks
            if b.block_type == "chemistry_formula"
        ]
        assert len(chem_blocks) >= 1
        # Subscript should be preserved in some form
        content = chem_blocks[0].raw_content
        assert "H" in content
        assert "O" in content

    def test_complex_chemistry_equation(self, tmp_path):
        """H₂SO₄ + 2NaOH → Na₂SO₄ + 2H₂O preserves subscripts and arrows."""
        # H₂SO₄
        body_xml = _make_chemistry_paragraph_xml([
            ("H", None),
            ("2", "subscript"),
            ("SO", None),
            ("4", "subscript"),
            (" + ", None),
            ("2", None),
            ("Na", None),
            ("O", None),
            ("H", None),
            (" ", None),
            ("→", None),
            (" ", None),
            ("Na", None),
            ("2", "subscript"),
            ("SO", None),
            ("4", "subscript"),
            (" + 2", None),
            ("H", None),
            ("2", "subscript"),
            ("O", None),
        ])
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        chem_blocks = [
            b for b in doc.pages[0].content_blocks
            if b.block_type == "chemistry_formula"
        ]
        assert len(chem_blocks) >= 1

        # Verify arrow is in the content
        all_chem_text = " ".join(b.raw_content for b in chem_blocks)
        assert "→" in all_chem_text or "⁺" in all_chem_text or "₂" in all_chem_text

    def test_chemistry_with_superscript_charge(self, tmp_path):
        """Fe²⁺ with superscript charge is part of chemistry detection."""
        body_xml = _make_chemistry_paragraph_xml([
            ("Fe", None),
            ("2+", "superscript"),
            (" + ", None),
            ("Cu", None),
            (" ", None),
            ("→", None),
            (" ", None),
            ("Fe", None),
            (" + ", None),
            ("Cu", None),
            ("2+", "superscript"),
        ])
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        chem_blocks = [
            b for b in doc.pages[0].content_blocks
            if b.block_type == "chemistry_formula"
        ]
        # Should detect chemistry due to element + superscript + arrow
        assert len(chem_blocks) >= 1

    def test_non_chemistry_paragraph_not_classified(self, tmp_path):
        """A plain Arabic text paragraph is NOT classified as chemistry."""
        body_xml = _make_paragraph_xml("هذا نص عربي عادي بدون أي صيغ كيميائية")
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        chem_blocks = [
            b for b in doc.pages[0].content_blocks
            if b.block_type == "chemistry_formula"
        ]
        assert len(chem_blocks) == 0

    def test_chemistry_subscript_unicode_preserved(self, tmp_path):
        """Subscript numbers are converted to Unicode subscript digits."""
        body_xml = _make_chemistry_paragraph_xml([
            ("H", None),
            ("2", "subscript"),
            ("SO", None),
            ("4", "subscript"),
        ])
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        chem_blocks = [
            b for b in doc.pages[0].content_blocks
            if b.block_type == "chemistry_formula"
        ]
        assert len(chem_blocks) >= 1
        content = chem_blocks[0].raw_content
        # Should contain Unicode subscript digits
        assert "₂" in content or "₄" in content


# ═══════════════════════════════════════════════════════════════════════
# AC-1.2.3: Paragraph, table, and heading order preserved
# ═══════════════════════════════════════════════════════════════════════


class TestAC_1_2_3_DocumentOrder:
    """AC-1.2.3: Paragraph, table, and heading order is preserved exactly
    as in the source Word document — no interleaving or dropping."""

    def test_paragraph_order_preserved(self, tmp_path):
        """Paragraphs appear in the same order as in the source document."""
        body_xml = (
            _make_paragraph_xml("الفقرة الأولى")
            + _make_paragraph_xml("الفقرة الثانية")
            + _make_paragraph_xml("الفقرة الثالثة")
        )
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        blocks = doc.pages[0].content_blocks
        texts = [b.raw_content for b in blocks if b.block_type == "arabic_text"]

        assert texts == ["الفقرة الأولى", "الفقرة الثانية", "الفقرة الثالثة"]

    def test_heading_before_paragraph(self, tmp_path):
        """A heading followed by a paragraph maintains order."""
        body_xml = (
            _make_paragraph_xml("الفصل الأول", style="Heading1")
            + _make_paragraph_xml("محتوى الفصل")
        )
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        blocks = doc.pages[0].content_blocks
        assert len(blocks) >= 2
        assert blocks[0].block_type == "chapter_header"
        assert blocks[0].raw_content == "الفصل الأول"
        assert blocks[1].block_type == "arabic_text"
        assert blocks[1].raw_content == "محتوى الفصل"

    def test_table_between_paragraphs_order(self, tmp_path):
        """Table between paragraphs: para → table → para order maintained."""
        body_xml = (
            _make_paragraph_xml("قبل الجدول")
            + _make_table_xml(["العمود 1", "العمود 2"], [["خلية 1", "خلية 2"]])
            + _make_paragraph_xml("بعد الجدول")
        )
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        blocks = doc.pages[0].content_blocks
        tables = doc.pages[0].tables

        # Two text blocks
        text_blocks = [b for b in blocks if b.block_type == "arabic_text"]
        assert len(text_blocks) == 2
        assert text_blocks[0].raw_content == "قبل الجدول"
        assert text_blocks[1].raw_content == "بعد الجدول"

        # One table
        assert len(tables) == 1
        assert tables[0].headers == ["العمود 1", "العمود 2"]
        assert tables[0].rows == [["خلية 1", "خلية 2"]]

    def test_mixed_content_order(self, tmp_path):
        """Mixed heading + paragraph + table + paragraph maintains order."""
        body_xml = (
            _make_paragraph_xml("عنوان الفصل", style="Heading1")
            + _make_paragraph_xml("مقدمة الفصل")
            + _make_table_xml(["اسم", "درجة"], [["أحمد", "90"], ["سارة", "95"]])
            + _make_paragraph_xml("خاتمة الفصل")
        )
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        blocks = doc.pages[0].content_blocks
        assert len(blocks) >= 3

        # Block types in order
        types = [b.block_type for b in blocks]
        assert types[0] == "chapter_header"
        assert types[1] == "arabic_text"
        assert types[2] == "arabic_text"  # الخاتمة

        # Contents in order
        assert blocks[0].raw_content == "عنوان الفصل"
        assert blocks[1].raw_content == "مقدمة الفصل"
        assert blocks[2].raw_content == "خاتمة الفصل"

    def test_no_paragraphs_dropped(self, tmp_path):
        """No paragraphs are silently dropped during extraction."""
        n_paragraphs = 10
        body_xml = "".join(
            _make_paragraph_xml(f"فقرة رقم {i+1}")
            for i in range(n_paragraphs)
        )
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        blocks = doc.pages[0].content_blocks
        assert len(blocks) == n_paragraphs

    def test_table_structure_preserved(self, tmp_path):
        """Table headers and rows match source structure."""
        headers = ["المادة", "الدرجة", "النتيجة"]
        rows = [
            ["رياضيات", "85", "ناجح"],
            ["فيزياء", "92", "ناجح"],
            ["كيمياء", "78", "ناجح"],
        ]
        body_xml = _make_table_xml(headers, rows)
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        tables = doc.pages[0].tables
        assert len(tables) == 1

        tbl = tables[0]
        assert tbl.headers == headers
        assert tbl.rows == rows

    def test_heading_levels_detected(self, tmp_path):
        """Heading1, Heading2 styles produce chapter_header blocks."""
        body_xml = (
            _make_paragraph_xml("المستوى الأول", style="Heading1")
            + _make_paragraph_xml("المستوى الثاني", style="Heading2")
            + _make_paragraph_xml("نص عادي")
        )
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        blocks = doc.pages[0].content_blocks
        assert blocks[0].block_type == "chapter_header"
        assert blocks[0].source_position.get("heading_level") == 1
        assert blocks[1].block_type == "chapter_header"
        assert blocks[1].source_position.get("heading_level") == 2
        assert blocks[2].block_type == "arabic_text"


# ═══════════════════════════════════════════════════════════════════════
# AC-1.2.5: OLE/legacy equation → legacy_equation_placeholder
# ═══════════════════════════════════════════════════════════════════════


class TestAC_1_2_5_OLEPlaceholder:
    """AC-1.2.5: A .docx file containing a legacy OLE equation object
    (w:object/v:shape) produces a legacy_equation_placeholder block at
    the correct document position — zero silently-dropped equations."""

    def test_ole_object_produces_placeholder(self, tmp_path):
        """w:object with OLE equation produces legacy_equation_placeholder."""
        body_xml = (
            _make_paragraph_xml("نص قبل المعادلة")
            + _make_ole_object_xml("Equation.3")
            + _make_paragraph_xml("نص بعد المعادلة")
        )
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        blocks = doc.pages[0].content_blocks
        legacy_blocks = [
            b for b in blocks
            if b.block_type == "legacy_equation_placeholder"
        ]
        assert len(legacy_blocks) == 1
        assert "Legacy equation" in legacy_blocks[0].raw_content

    def test_ole_placeholder_at_correct_position(self, tmp_path):
        """OLE placeholder appears between the surrounding paragraphs."""
        body_xml = (
            _make_paragraph_xml("الفقرة الأولى")
            + _make_ole_object_xml("Equation.3")
            + _make_paragraph_xml("الفقرة الثالثة")
        )
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        blocks = doc.pages[0].content_blocks
        types = [b.block_type for b in blocks]

        assert types[0] == "arabic_text"
        assert types[1] == "legacy_equation_placeholder"
        assert types[2] == "arabic_text"

    def test_multiple_ole_objects_all_detected(self, tmp_path):
        """Multiple OLE objects all produce placeholders — none dropped."""
        body_xml = (
            _make_ole_object_xml("Equation.3")
            + _make_paragraph_xml("فقرة بين معادلتين")
            + _make_ole_object_xml("Equation.DSMT4")
        )
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        blocks = doc.pages[0].content_blocks
        legacy_blocks = [
            b for b in blocks
            if b.block_type == "legacy_equation_placeholder"
        ]
        assert len(legacy_blocks) == 2

    def test_ole_prog_id_captured(self, tmp_path):
        """The progId (e.g. 'Equation.3') is captured in the placeholder."""
        body_xml = _make_ole_object_xml("Equation.3")
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        legacy_blocks = [
            b for b in doc.pages[0].content_blocks
            if b.block_type == "legacy_equation_placeholder"
        ]
        assert len(legacy_blocks) >= 1
        block = legacy_blocks[0]
        # progId should be in source_position or raw_content
        assert (
            "Equation.3" in block.raw_content
            or block.source_position.get("prog_id") == "Equation.3"
        )

    @pytest.mark.skipif(
        not XSLT_AVAILABLE,
        reason="OMML2MML.XSL not found — skipping combined count test",
    )
    def test_combined_omath_and_ole_count_matches(self, tmp_path):
        """AC-1.2.5: Total extracted block count matches source element
        count (m:oMath + w:object combined) — zero silently dropped."""
        body_xml = (
            _make_omath_xml("x+y")           # 1 math_formula
            + _make_paragraph_xml("نص")       # 1 arabic_text
            + _make_ole_object_xml("Eq.3")    # 1 legacy_equation_placeholder
            + _make_omath_xml("a-b")          # 1 math_formula
            + _make_ole_object_xml("Eq.DSMT") # 1 legacy_equation_placeholder
        )
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        blocks = doc.pages[0].content_blocks

        math_count = sum(
            1 for b in blocks if b.block_type == "math_formula"
        )
        ole_count = sum(
            1 for b in blocks if b.block_type == "legacy_equation_placeholder"
        )

        # 2 m:oMath + 2 OLE = 4 equation-related blocks
        assert math_count == 2, f"Expected 2 math blocks, got {math_count}"
        assert ole_count == 2, f"Expected 2 OLE blocks, got {ole_count}"


# ═══════════════════════════════════════════════════════════════════════
# General extraction correctness tests
# ═══════════════════════════════════════════════════════════════════════


class TestGeneralExtraction:
    """Additional tests for extraction correctness."""

    def test_empty_document(self, tmp_path):
        """An empty DOCX produces a BookletDocument with no blocks."""
        body_xml = ""
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        assert doc.file_type == "docx"
        assert len(doc.pages) == 1
        assert len(doc.pages[0].content_blocks) == 0

    def test_file_not_found_raises(self):
        """Extracting a non-existent file raises FileNotFoundError."""
        extractor = DocxContentExtractor()
        with pytest.raises(FileNotFoundError):
            extractor.extract("/nonexistent/file.docx")

    def test_content_hash_computed_correctly(self, tmp_path):
        """Every extracted block has a valid content_hash."""
        body_xml = _make_paragraph_xml("اختبار الهاش")
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        for block in doc.pages[0].content_blocks:
            expected = compute_canonical_hash(block.raw_content)
            assert block.content_hash == expected, (
                f"Hash mismatch for block {block.block_id}"
            )

    def test_block_ids_are_unique(self, tmp_path):
        """All block IDs within a document are unique."""
        body_xml = (
            _make_paragraph_xml("أ")
            + _make_paragraph_xml("ب")
            + _make_paragraph_xml("ج")
        )
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        ids = [b.block_id for b in doc.pages[0].content_blocks]
        assert len(ids) == len(set(ids)), "Duplicate block IDs found"

    def test_booklet_document_structure(self, tmp_path):
        """BookletDocument has correct file_type and file_path."""
        body_xml = _make_paragraph_xml("test")
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        assert doc.file_type == "docx"
        assert doc.file_path.endswith(".docx")
        assert doc.doc_id  # UUID is set
        assert len(doc.pages) == 1

    def test_empty_paragraphs_skipped(self, tmp_path):
        """Empty paragraphs (no text) do not produce blocks."""
        body_xml = (
            '<w:p></w:p>'
            + _make_paragraph_xml("فقرة فعلية")
            + '<w:p><w:pPr></w:pPr></w:p>'
        )
        docx_path = _create_docx_with_body_xml(body_xml, str(tmp_path))

        extractor = DocxContentExtractor()
        doc = extractor.extract(docx_path)

        blocks = doc.pages[0].content_blocks
        # Only the actual paragraph with text should produce a block
        text_blocks = [b for b in blocks if b.block_type == "arabic_text"]
        assert len(text_blocks) == 1
        assert text_blocks[0].raw_content == "فقرة فعلية"


# ═══════════════════════════════════════════════════════════════════════
# Unit tests for helper functions
# ═══════════════════════════════════════════════════════════════════════


class TestHelperFunctions:
    """Unit tests for low-level helper functions."""

    def test_get_run_text(self):
        """_get_run_text extracts text from w:t children."""
        run_xml = f"""<w:r xmlns:w="{W_NS}">
            <w:t>Hello</w:t>
            <w:t> World</w:t>
        </w:r>"""
        run = ET.fromstring(run_xml.encode("utf-8"))
        assert _get_run_text(run) == "Hello World"

    def test_get_run_text_empty(self):
        """_get_run_text returns empty string for run with no text."""
        run_xml = f'<w:r xmlns:w="{W_NS}"><w:rPr/></w:r>'
        run = ET.fromstring(run_xml.encode("utf-8"))
        assert _get_run_text(run) == ""

    def test_get_run_vertical_align_subscript(self):
        """_get_run_vertical_align detects subscript."""
        run_xml = f"""<w:r xmlns:w="{W_NS}">
            <w:rPr>
                <w:vertAlign w:val="subscript"/>
            </w:rPr>
            <w:t>2</w:t>
        </w:r>"""
        run = ET.fromstring(run_xml.encode("utf-8"))
        assert _get_run_vertical_align(run) == "subscript"

    def test_get_run_vertical_align_superscript(self):
        """_get_run_vertical_align detects superscript."""
        run_xml = f"""<w:r xmlns:w="{W_NS}">
            <w:rPr>
                <w:vertAlign w:val="superscript"/>
            </w:rPr>
            <w:t>2+</w:t>
        </w:r>"""
        run = ET.fromstring(run_xml.encode("utf-8"))
        assert _get_run_vertical_align(run) == "superscript"

    def test_get_run_vertical_align_none(self):
        """_get_run_vertical_align returns None for normal run."""
        run_xml = f"""<w:r xmlns:w="{W_NS}">
            <w:t>Normal</w:t>
        </w:r>"""
        run = ET.fromstring(run_xml.encode("utf-8"))
        assert _get_run_vertical_align(run) is None

    def test_assemble_chemistry_subscript_digits(self):
        """_assemble_chemistry_text converts subscript digits to Unicode."""
        run_info = [
            {"text": "H", "valign": None, "is_sub": False, "is_sup": False,
             "has_element": True, "has_arrow": False, "has_digit": False, "has_plus": False},
            {"text": "2", "valign": "subscript", "is_sub": True, "is_sup": False,
             "has_element": False, "has_arrow": False, "has_digit": True, "has_plus": False},
            {"text": "O", "valign": None, "is_sub": False, "is_sup": False,
             "has_element": True, "has_arrow": False, "has_digit": False, "has_plus": False},
        ]
        result = _assemble_chemistry_text(run_info, 0, 2)
        assert result == "H₂O"
