"""
Booklet_AI Desktop — Tests for STEP-7.6

Validates that image blocks are not treated as text:
- (a) A BookletDocument with text blocks + one image_prompt block (with
      source_position image_asset) passes the coverage audit with text_loss == False.
- (b) The same document with one text block genuinely removed still reports
      text_loss == True (the audit must not become blind).
- (c) RuleBasedClassifier leaves an image_prompt block with Latin characters as image_prompt.
- (d) OpenRouterClient.classify_blocks with a mocked HTTP response does not include
      the image block id in the payload and still passes the exact-match validation.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import httpx
import pytest
from docx import Document

from app.ai.classifier import PRECLASSIFIED_TYPES, RuleBasedClassifier
from app.ai.client import OpenRouterClient
from app.core.models import BookletDocument, ContentBlock, PageData
from app.core.text_integrity import compute_canonical_hash
from tools.generate_from_docx import _is_non_text_block, run_coverage_audit


# ─── Helper Functions ──────────────────────────────────────────────────


def _make_block(
    block_id: str,
    raw_content: str,
    block_type: str = "arabic_text",
    source_position: dict[str, Any] | None = None,
) -> ContentBlock:
    """Helper to create a ContentBlock with matching canonical hash."""
    return ContentBlock(
        block_id=block_id,
        block_type=block_type,
        raw_content=raw_content,
        content_hash=compute_canonical_hash(raw_content),
        source_position=source_position or {},
    )


# ─── Test Cases ────────────────────────────────────────────────────────


class TestStep76ImageBlocksNotTreatedAsText:
    """Acceptance test suite for STEP-7.6."""

    def test_coverage_audit_with_image_prompt_block_passes(self, tmp_path: Path):
        """(a) a BookletDocument with text blocks + one image_prompt block
        (with source_position image_asset) passes coverage audit with text_loss == False."""
        docx_path = tmp_path / "synthetic_audit.docx"
        doc = Document()
        doc.add_paragraph("المقدمة الأولى في الكتاب")
        doc.add_paragraph("الفصل الأول يحتوي على شرح تفصيلي")
        doc.save(str(docx_path))

        b1 = _make_block("b1", "المقدمة الأولى في الكتاب", block_type="arabic_text")
        b2 = _make_block(
            "b2", "الفصل الأول يحتوي على شرح تفصيلي", block_type="arabic_text"
        )
        b_img = _make_block(
            "b_img",
            "[image:docx_img_001.png]",
            block_type="image_prompt",
            source_position={
                "image_asset": {
                    "source_path": "docx_img_001.png",
                    "pixel_width": 400,
                    "pixel_height": 300,
                }
            },
        )
        page = PageData(
            page_number=1,
            content_blocks=[b1, b2, b_img],
            tables=[],
            width=210.0,
            height=297.0,
        )
        booklet_doc = BookletDocument(
            doc_id="synthetic_doc",
            file_path=str(docx_path),
            file_type="docx",
            source_sha256="dummy_sha256",
            pages=[page],
            audit_rows=[],
        )

        audit_rows, audit_warnings, text_loss = run_coverage_audit(
            docx_path, booklet_doc
        )

        assert text_loss is False
        assert "TEXT_DROPPED" not in audit_warnings

    def test_coverage_audit_with_dropped_text_reports_loss(self, tmp_path: Path):
        """(b) the same document with one text block genuinely removed still
        reports text_loss == True (the audit must not become blind)."""
        docx_path = tmp_path / "synthetic_audit_missing.docx"
        doc = Document()
        doc.add_paragraph("المقدمة الأولى في الكتاب")
        doc.add_paragraph("الفصل الأول يحتوي على شرح تفصيلي")
        doc.save(str(docx_path))

        b1 = _make_block("b1", "المقدمة الأولى في الكتاب", block_type="arabic_text")
        # b2 ("الفصل الأول يحتوي على شرح تفصيلي") is genuinely omitted
        b_img = _make_block(
            "b_img",
            "[image:docx_img_001.png]",
            block_type="image_prompt",
            source_position={
                "image_asset": {
                    "source_path": "docx_img_001.png",
                    "pixel_width": 400,
                    "pixel_height": 300,
                }
            },
        )
        page = PageData(
            page_number=1,
            content_blocks=[b1, b_img],
            tables=[],
            width=210.0,
            height=297.0,
        )
        booklet_doc = BookletDocument(
            doc_id="synthetic_doc_missing",
            file_path=str(docx_path),
            file_type="docx",
            source_sha256="dummy_sha256",
            pages=[page],
            audit_rows=[],
        )

        audit_rows, audit_warnings, text_loss = run_coverage_audit(
            docx_path, booklet_doc
        )

        assert text_loss is True
        assert "TEXT_DROPPED" in audit_warnings

    def test_rule_based_classifier_preserves_image_prompt_with_latin(self):
        """(c) RuleBasedClassifier leaves an image_prompt block with Latin characters
        as image_prompt."""
        raw_text = "[image:docx_img_01_latin_chars.png]"
        block = _make_block(
            "b_img",
            raw_text,
            block_type="image_prompt",
            source_position={
                "image_asset": {
                    "source_path": "docx_img_01_latin_chars.png",
                    "pixel_width": 300,
                    "pixel_height": 200,
                }
            },
        )

        # Single block classification
        classified_type = RuleBasedClassifier.classify_block(block)
        assert classified_type == "image_prompt"

        # Document-level batch classification
        classified_blocks = RuleBasedClassifier.classify_document_blocks([block])
        assert len(classified_blocks) == 1
        assert classified_blocks[0].block_type == "image_prompt"
        assert classified_blocks[0].raw_content == raw_text
        assert classified_blocks[0].content_hash == block.content_hash

    def test_openrouter_client_excludes_image_prompt_and_validates_eligible(self):
        """(d) OpenRouterClient.classify_blocks with a mocked HTTP response
        does not include the image block id in the payload and still passes
        the exact-match validation."""
        text_content = "هذا نص تجريبي للاختبار"
        b_text = _make_block(
            "b_text_1",
            text_content,
            block_type="arabic_text",
        )
        b_img = _make_block(
            "b_img_1",
            "[image:docx_img_01.png]",
            block_type="image_prompt",
            source_position={"image_asset": {"source_path": "docx_img_01.png"}},
        )

        client = OpenRouterClient(api_key="sk-synthetic-test-key")

        captured_payloads = []

        def fake_send_request(payload: dict) -> httpx.Response:
            captured_payloads.append(payload)
            user_content = payload["messages"][1]["content"]
            items = json.loads(user_content)
            req_bids = [item["block_id"] for item in items]

            # The payload must NEVER include the image block id
            assert "b_img_1" not in req_bids
            assert "b_text_1" in req_bids

            # Mock LLM returns classifications only for requested eligible ids
            response_body = {
                "classifications": [
                    {"block_id": bid, "type": "arabic_text"}
                    for bid in req_bids
                ]
            }
            return httpx.Response(
                status_code=200,
                json={
                    "choices": [
                        {
                            "message": {
                                "content": json.dumps(response_body),
                            }
                        }
                    ]
                },
                request=httpx.Request(
                    "POST", "https://openrouter.ai/api/v1/chat/completions"
                ),
            )

        client._send_request = fake_send_request

        results = client.classify_blocks([b_text, b_img])

        assert len(captured_payloads) == 1
        assert len(results) == 2

        # Text block was classified and hashes preserved
        assert results[0].block_id == "b_text_1"
        assert results[0].block_type == "arabic_text"
        assert results[0].raw_content == text_content
        assert results[0].content_hash == b_text.content_hash

        # Image block retained its image_prompt type and hashes preserved
        assert results[1].block_id == "b_img_1"
        assert results[1].block_type == "image_prompt"
        assert results[1].raw_content == "[image:docx_img_01.png]"
        assert results[1].content_hash == b_img.content_hash

    def test_is_non_text_block_helper(self):
        """Helper correctly identifies image_prompt, legacy equations, formulas,
        and blocks with source_position['image_asset']."""
        # Math formula
        b_math = _make_block("m1", "x + y = z", block_type="math_formula")
        assert _is_non_text_block(b_math) is True

        # Legacy equation placeholder
        b_ole = _make_block("ole1", "[OLE Equation]", block_type="legacy_equation_placeholder")
        assert _is_non_text_block(b_ole) is True

        # Image prompt
        b_img = _make_block("img1", "[image:a.png]", block_type="image_prompt")
        assert _is_non_text_block(b_img) is True

        # Arbitrary block with image_asset in source_position
        b_sp_img = _make_block(
            "custom1",
            "some text",
            block_type="arabic_text",
            source_position={"image_asset": {"source_path": "a.png"}},
        )
        assert _is_non_text_block(b_sp_img) is True

        # Normal text blocks are NOT non-text
        b_ar = _make_block("ar1", "نص عربي", block_type="arabic_text")
        assert _is_non_text_block(b_ar) is False

        b_en = _make_block("en1", "English text", block_type="english_text")
        assert _is_non_text_block(b_en) is False

    def test_preclassified_types_constants(self):
        """PRECLASSIFIED_TYPES contains all four pre-classified types."""
        assert "math_formula" in PRECLASSIFIED_TYPES
        assert "chemistry_formula" in PRECLASSIFIED_TYPES
        assert "image_prompt" in PRECLASSIFIED_TYPES
        assert "legacy_equation_placeholder" in PRECLASSIFIED_TYPES
