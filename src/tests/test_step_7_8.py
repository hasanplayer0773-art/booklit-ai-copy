"""
Tests for STEP-7.8 — Chemistry runs must never drop surrounding text (zero-mutation bug).

Verifies:
1. Exactly one block per paragraph (per run group) preserving ALL text in source order.
2. Normalized equality between source paragraph text and extracted block text (no text dropped).
3. Pure formulas become 'chemistry_formula'; sentences with inline formulas become 'arabic_text'.
4. run_coverage_audit reports no text loss (mapping U+2212 to '-' on both sides).
"""

from __future__ import annotations

import unicodedata
from pathlib import Path

import pytest
from docx import Document

from app.core.docx_extractor import DocxContentExtractor
from app.core.text_integrity import verify_text_integrity
from tools.generate_from_docx import run_coverage_audit


def _audit_norm(text: str) -> str:
    """Normalize text with NFKC and map U+2212 minus sign to '-' for comparison."""
    return unicodedata.normalize("NFKC", text).replace("\u2212", "-")


class TestStep78ChemistryZeroMutation:
    """STEP-7.8: Zero-mutation guarantee for chemistry runs and surrounding text."""

    def test_step_7_8_full_specification(self, tmp_path: Path):
        """Construct synthetic docx with the 6 required paragraphs and verify:
        (a) exactly one block per paragraph
        (b) NFKC(source paragraph text, with U+2212->'-') == NFKC(block raw_content, with U+2212->'-')
        (c) pure-formula is chemistry_formula, others are arabic_text
        (d) run_coverage_audit reports no text loss
        """
        doc = Document()

        # Paragraph 1: pure formula
        p1 = doc.add_paragraph()
        p1.add_run("H")
        p1.add_run("2").font.subscript = True
        p1.add_run("SO")
        p1.add_run("4").font.subscript = True
        p1.add_run(" + 2NaOH → Na")
        p1.add_run("2").font.subscript = True
        p1.add_run("SO")
        p1.add_run("4").font.subscript = True
        p1.add_run(" + 2H")
        p1.add_run("2").font.subscript = True
        p1.add_run("O")

        # Paragraph 2: "يتفاعل " + H,2sub,SO,4sub
        p2 = doc.add_paragraph()
        p2.add_run("يتفاعل ")
        p2.add_run("H")
        p2.add_run("2").font.subscript = True
        p2.add_run("SO")
        p2.add_run("4").font.subscript = True

        # Paragraph 3: "يتفاعل " + formula + " مع هيدروكسيد الصوديوم ويعطي ملحاً وماءً"
        p3 = doc.add_paragraph()
        p3.add_run("يتفاعل ")
        p3.add_run("H")
        p3.add_run("2").font.subscript = True
        p3.add_run("SO")
        p3.add_run("4").font.subscript = True
        p3.add_run(" مع هيدروكسيد الصوديوم ويعطي ملحاً وماءً")

        # Paragraph 4: formula + " هو الماء وصيغته معروفة عند الطلاب"
        p4 = doc.add_paragraph()
        p4.add_run("H")
        p4.add_run("2").font.subscript = True
        p4.add_run("O")
        p4.add_run(" هو الماء وصيغته معروفة عند الطلاب")

        # Paragraph 5: "الأيونات الناتجة هي Na" + "+" sup + " و SO" + "4" sub + "2-" sup + " في المحلول"
        p5 = doc.add_paragraph()
        p5.add_run("الأيونات الناتجة هي Na")
        p5.add_run("+").font.superscript = True
        p5.add_run(" و SO")
        p5.add_run("4").font.subscript = True
        p5.add_run("2-").font.superscript = True
        p5.add_run(" في المحلول")

        # Paragraph 6: a plain Arabic sentence
        p6 = doc.add_paragraph("هذه فقرة نصية عادية باللغة العربية لا تحتوي على صيغ كيميائية.")

        docx_path = tmp_path / "step_7_8_test.docx"
        doc.save(str(docx_path))

        extractor = DocxContentExtractor()
        booklet_doc = extractor.extract(str(docx_path))
        blocks = booklet_doc.pages[0].content_blocks

        # (a) Exactly one block per paragraph
        assert len(blocks) == len(doc.paragraphs) == 6, (
            f"Expected exactly 6 blocks for 6 paragraphs, got {len(blocks)}"
        )

        # (b) For every paragraph: NFKC(source, U+2212->'-') == NFKC(block raw_content, U+2212->'-')
        for i, (para, block) in enumerate(zip(doc.paragraphs, blocks)):
            src_norm = _audit_norm(para.text)
            blk_norm = _audit_norm(block.raw_content)
            assert src_norm == blk_norm, (
                f"Paragraph {i} text mismatch: source={repr(src_norm)} vs block={repr(blk_norm)}"
            )
            # content_hash must match final raw_content
            verify_text_integrity(block.raw_content, block.content_hash)

        # (c) Block types: pure formula is chemistry_formula, others are arabic_text
        assert blocks[0].block_type == "chemistry_formula"
        assert blocks[0].raw_content == "H₂SO₄ + 2NaOH → Na₂SO₄ + 2H₂O"

        assert blocks[1].block_type == "arabic_text"
        assert blocks[1].raw_content == "يتفاعل H₂SO₄"

        assert blocks[2].block_type == "arabic_text"
        assert blocks[2].raw_content == "يتفاعل H₂SO₄ مع هيدروكسيد الصوديوم ويعطي ملحاً وماءً"

        assert blocks[3].block_type == "arabic_text"
        assert blocks[3].raw_content == "H₂O هو الماء وصيغته معروفة عند الطلاب"

        assert blocks[4].block_type == "arabic_text"
        assert blocks[4].raw_content == "الأيونات الناتجة هي Na⁺ و SO₄²⁻ في المحلول"

        assert blocks[5].block_type == "arabic_text"
        assert blocks[5].raw_content == "هذه فقرة نصية عادية باللغة العربية لا تحتوي على صيغ كيميائية."

        # (d) run_coverage_audit on this document reports no text loss
        audit_rows, audit_warnings, text_loss = run_coverage_audit(str(docx_path), booklet_doc)
        assert not text_loss, "run_coverage_audit reported text loss"
        assert "TEXT_DROPPED" not in audit_warnings, f"Unexpected audit warnings: {audit_warnings}"

    def test_step_7_8_standalone_pure_formula(self, tmp_path: Path):
        """Single compound with subscripts (e.g. H₂SO₄) produces chemistry_formula."""
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("H")
        p.add_run("2").font.subscript = True
        p.add_run("SO")
        p.add_run("4").font.subscript = True

        docx_path = tmp_path / "standalone_formula.docx"
        doc.save(str(docx_path))

        extractor = DocxContentExtractor()
        booklet_doc = extractor.extract(str(docx_path))
        blocks = booklet_doc.pages[0].content_blocks

        assert len(blocks) == 1
        assert blocks[0].block_type == "chemistry_formula"
        assert blocks[0].raw_content == "H₂SO₄"
        assert blocks[0].source_position.get("type") == "chemistry_runs"
        verify_text_integrity(blocks[0].raw_content, blocks[0].content_hash)

    def test_step_7_8_pure_formula_with_punctuation(self, tmp_path: Path):
        """Pure formula paragraph with trailing punctuation or parentheses remains chemistry_formula."""
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("(")
        p.add_run("H")
        p.add_run("2").font.subscript = True
        p.add_run("O")
        p.add_run(").")

        docx_path = tmp_path / "punct_formula.docx"
        doc.save(str(docx_path))

        extractor = DocxContentExtractor()
        booklet_doc = extractor.extract(str(docx_path))
        blocks = booklet_doc.pages[0].content_blocks

        assert len(blocks) == 1
        assert blocks[0].block_type == "chemistry_formula"
        assert blocks[0].raw_content == "(H₂O)."

    def test_step_7_8_heading_preserves_heading_level(self, tmp_path: Path):
        """A heading style applied to text containing chemistry becomes chapter_header."""
        doc = Document()
        p = doc.add_heading(level=2)
        p.add_run("خواص حمض ")
        p.add_run("H")
        p.add_run("2").font.subscript = True
        p.add_run("SO")
        p.add_run("4").font.subscript = True

        docx_path = tmp_path / "heading_chem.docx"
        doc.save(str(docx_path))

        extractor = DocxContentExtractor()
        booklet_doc = extractor.extract(str(docx_path))
        blocks = booklet_doc.pages[0].content_blocks

        assert len(blocks) == 1
        assert blocks[0].block_type == "chapter_header"
        assert blocks[0].source_position.get("heading_level") == 2
        assert blocks[0].raw_content == "خواص حمض H₂SO₄"
